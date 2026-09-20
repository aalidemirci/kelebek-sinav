"""BEP kapsamındaki öğrenciler + bireysel soru dosyası testleri (20.09.2026).

Kullanıcı isteğinin üç güvencesi burada sabitlenir:
1. Seçilen öğrencinin kitapçığı KENDİ PDF'inden, ADINA basılır.
2. Öğrenciyi AYIRAN HİÇBİR İŞARET salon evrakına, duyuruya, çizelgeye, tutanağa,
   doğrulama raporuna ve kitapçık bandına basılmaz; idare özeti pakete girmez.
3. KVKK md. 6: öğrenci bağı şifrelidir, satırlar KATI silinir, hata/günlük
   metinlerinde ve URL yollarında öğrenci kimliği geçmez.
"""

from __future__ import annotations

import io
import logging
import re
import zipfile
from typing import Any

import pytest
from django.core.exceptions import ValidationError
from django.db import connection
from openpyxl import load_workbook
from pypdf import PdfReader
from rest_framework.test import APIClient

from apps.okul.models import Student, StudentStatus
from apps.okul.services import persons
from apps.sinav import booklet, services, services_individual
from apps.sinav.models import (
    BookletRunStatus,
    ExamSession,
    ExamSessionStatus,
    IepStudent,
    IndividualQuestionDocument,
    ParticipantType,
    ScoreMode,
    SeatAssignment,
)
from apps.sinav.serializers import BookletRunSerializer
from apps.sinav.tests.oturum_yardim import ders, oturum, salon, sube
from apps.sinav.tests.test_booklets import PLAN_3X2_DOUBLE, _question_pdf
from shared import crypto

pytestmark = pytest.mark.django_db

#: Bireysel PDF'in gövdesinde aranacak işaret — kitapçığın KAYNAĞINI kanıtlar.
BIREYSEL_IZ = "UYARLANMIS-SORULAR"


def _session(*, with_course_docs: bool = True) -> ExamSession:
    """2 ders × 4 öğrenci, tek salon, dağıtılmış; istenirse ders dosyaları yüklü."""
    sube(9, "A", students=4, start_no=101)
    sube(10, "A", students=4, start_no=201)
    c9 = ders("Coğrafya", levels=[9])
    c10 = ders("Fizik", levels=[10])
    session = oturum()
    sc9 = services.add_session_course(
        session, course_id=c9.pk, participant_type=ParticipantType.LEVEL, level=9
    )
    sc10 = services.add_session_course(
        session, course_id=c10.pk, participant_type=ParticipantType.LEVEL, level=10
    )
    services.set_session_rooms(session, [{"room_id": salon("D-201", plan=PLAN_3X2_DOUBLE).pk}])
    services.distribute_session(session, seed=42)
    if with_course_docs:
        services.upload_question_document(sc9, file_bytes=_question_pdf(2, title="Coğrafya"))
        services.upload_question_document(sc10, file_bytes=_question_pdf(2, title="Fizik"))
    session.refresh_from_db()
    return session


def _student(number: str) -> Student:
    student: Student = Student.objects.get(student_number=number)
    return student


def _select_with_file(session: ExamSession, number: str, *, pages: int = 1) -> Any:
    student = _student(number)
    if student.pk not in services_individual.iep_student_ids():
        services_individual.add_iep_student(student.pk)
    row = services_individual.select_individual(session, student_id=student.pk)
    return services_individual.upload_individual(
        row, file_bytes=_question_pdf(pages, title=BIREYSEL_IZ)
    )


def _pdf_text(pdf_bytes: bytes) -> str:
    """Sayfa metni, boşlukları teke indirilmiş (satır sonu sarması iddiayı kırmasın)."""
    pages = PdfReader(io.BytesIO(pdf_bytes)).pages
    return " ".join(" ".join(page.extract_text() or "" for page in pages).split())


def _room_pdf_pages(run: Any) -> list[str]:
    with zipfile.ZipFile(io.BytesIO(run.file.read())) as zf:
        data = zf.read(zf.namelist()[0])
    return [page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages]


# ===========================================================================
# BEP kapsamındaki öğrenciler listesi
# ===========================================================================


def test_liste_ekle_cikar_ve_sirala() -> None:
    sube(10, "A", students=2, start_no=201)
    sube(9, "A", students=2, start_no=101)
    for number in ("201", "102", "101"):
        services_individual.add_iep_student(_student(number).pk)

    items = services_individual.iep_list()
    # Sınıf/şube (seviye sayısal), sonra okul no — eklenme sırası DEĞİL.
    assert [item["student_number"] for item in items] == ["101", "102", "201"]
    assert set(items[0]) == {"id", "student_id", "student_number", "full_name", "class_label"}

    with pytest.raises(ValidationError, match="zaten listede"):
        services_individual.add_iep_student(_student("101").pk)

    services_individual.remove_iep_student(IepStudent.objects.get(pk=items[0]["id"]))
    assert [i["student_number"] for i in services_individual.iep_list()] == ["102", "201"]
    # Soft-delete YOK: çıkarılan kayıt iz bırakmaz.
    assert IepStudent.all_objects.count() == 2


def test_listeye_yalniz_aktif_ogrenci_girer() -> None:
    sube(9, "A", students=1, start_no=101)
    student = _student("101")
    Student.objects.filter(pk=student.pk).update(status=StudentStatus.LEFT)
    with pytest.raises(ValidationError, match="aktif değil"):
        services_individual.add_iep_student(student.pk)
    with pytest.raises(ValidationError, match="bulunamadı"):
        services_individual.add_iep_student(999_999)


def test_ogrenci_bagi_parola_acikken_sifreli_saklanir() -> None:
    """Okul no açık alandır; düz bağ "şu numaralı öğrenci BEP'li" demek olurdu."""
    sube(9, "A", students=1, start_no=101)
    student = _student("101")
    crypto.load_key(crypto.new_data_key())
    try:
        services_individual.add_iep_student(student.pk)
        with connection.cursor() as cursor:
            cursor.execute("SELECT student_ref FROM sinav_iepstudent")
            raw = str(cursor.fetchone()[0])
        assert raw != str(student.pk) and not raw.isdigit()  # Fernet token'ı
        assert services_individual.iep_student_ids() == {student.pk}
    finally:
        crypto.unload_key()
    # Kilitli kasa: bağ çözülemez → satır ATLANIR ama SİLİNMEZ (purge dokunmaz).
    assert services_individual.iep_student_ids() == set()
    assert services_individual.purge_stale() == 0
    assert IepStudent.objects.count() == 1


def test_sifreli_alanlar_parola_gecisine_kendiliginden_girer() -> None:
    from apps.okul.services import app_password

    covered = {model.__name__: fields for model, fields in app_password.encrypted_field_map()}
    assert covered["IepStudent"] == ("student_ref",)
    assert covered["IndividualQuestionDocument"] == ("student_ref",)


def test_modelde_serbest_metin_alani_yok() -> None:
    """KVKK md. 6: tanı/açıklama alanı EKLENMEZ — alan kümesi sabittir."""
    iep_fields = {f.name for f in IepStudent._meta.get_fields()}
    assert iep_fields == {"id", "created_at", "updated_at", "deleted_at", "student_ref"}
    doc_fields = {f.name for f in IndividualQuestionDocument._meta.get_fields()}
    assert doc_fields == {
        "id",
        "created_at",
        "updated_at",
        "deleted_at",
        "session",
        "student_ref",
        "file",
        "page_count",
        "sha256",
        "score_mode",
        "question_count",
    }


# ===========================================================================
# Seçim + yükleme
# ===========================================================================


def test_secim_kapilari() -> None:
    session = _session(with_course_docs=False)
    student = _student("101")

    with pytest.raises(ValidationError, match="listesindeki"):
        services_individual.select_individual(session, student_id=student.pk)

    services_individual.add_iep_student(student.pk)
    row = services_individual.select_individual(session, student_id=student.pk)
    assert not row.file and row.page_count is None
    with pytest.raises(ValidationError, match="zaten seçili"):
        services_individual.select_individual(session, student_id=student.pk)

    # Oturumda yerleşimi olmayan öğrenci seçilemez.
    sube(11, "A", students=1, start_no=301)
    outsider = _student("301")
    services_individual.add_iep_student(outsider.pk)
    with pytest.raises(ValidationError, match="yerleşiminde yok"):
        services_individual.select_individual(session, student_id=outsider.pk)


def test_yukleme_ders_dosyasiyla_ayni_dogrulamadan_gecer() -> None:
    session = _session(with_course_docs=False)
    student = _student("101")
    services_individual.add_iep_student(student.pk)
    row = services_individual.select_individual(session, student_id=student.pk)

    with pytest.raises(ValidationError, match="PDF değil"):
        services_individual.upload_individual(row, file_bytes=b"degil")
    with pytest.raises(ValidationError, match="soru sayısı"):
        services_individual.upload_individual(
            row, file_bytes=_question_pdf(1), score_mode=ScoreMode.QUESTION_TABLE
        )

    row = services_individual.upload_individual(
        row, file_bytes=_question_pdf(3), score_mode=ScoreMode.QUESTION_TABLE, question_count=6
    )
    assert row.page_count == 3 and row.question_count == 6 and len(row.sha256) == 64
    # Dosya adı öğrenciyi ANMAZ: rastgele 32 haneli onaltılık ad (pk, okul no, ad YOK).
    # Kalıp iddiası bilinçli: "101 geçmesin" demek rastgele ada denk gelip sahte kırmızı verir.
    name = row.file.name.rsplit("/", 1)[-1]
    assert re.fullmatch(r"soru_b_[0-9a-f]{32}[.]pdf", name), name


def test_degistirilen_ve_kaldirilan_dosya_diskten_silinir(
    django_capture_on_commit_callbacks: Any,
) -> None:
    session = _session(with_course_docs=False)
    first = _select_with_file(session, "101")
    storage, first_name = first.file.storage, first.file.name
    assert storage.exists(first_name)

    with django_capture_on_commit_callbacks(execute=True):
        second = services_individual.upload_individual(first, file_bytes=_question_pdf(2))
    assert not storage.exists(first_name), "değiştirilen eski dosya diskte kaldı"
    second_name = second.file.name
    assert storage.exists(second_name)

    with django_capture_on_commit_callbacks(execute=True):
        services_individual.remove_individual(second)
    assert not storage.exists(second_name), "kaldırılan dosya diskte kaldı"
    # KATI silme: satır iz bırakmaz.
    assert IndividualQuestionDocument.all_objects.count() == 0


def test_onayli_oturumda_degisiklik_reddedilir() -> None:
    session = _session()
    row = _select_with_file(session, "101")
    other = _student("102")
    services_individual.add_iep_student(other.pk)
    services.approve_session(session)
    session.refresh_from_db()
    row.refresh_from_db()

    with pytest.raises(ValidationError, match="değiştirilemez"):
        services_individual.select_individual(session, student_id=other.pk)
    with pytest.raises(ValidationError, match="değiştirilemez"):
        services_individual.upload_individual(row, file_bytes=_question_pdf(1))
    with pytest.raises(ValidationError, match="değiştirilemez"):
        services_individual.remove_individual(row)
    # Kitapçık yeniden basımı onaylı oturumda da bireysel dosyayı kullanır.
    run = services.request_booklet_run(session)
    assert run.status == BookletRunStatus.COMPLETED
    assert run.manifest["individual_booklets"] == 1


# ===========================================================================
# Kitapçık üretimi
# ===========================================================================


def test_kitapcik_ogrencinin_kendi_dosyasindan_adina_basilir() -> None:
    session = _session()
    _select_with_file(session, "101", pages=1)
    target = SeatAssignment.objects.get(session=session, student=_student("101"))

    run = services.request_booklet_run(session)
    assert run.status == BookletRunStatus.COMPLETED, run.error_message
    assert run.manifest["individual_booklets"] == 1
    # 7 öğrenci × 2 sayfa (ders dosyası) + 1 öğrenci × 1 sayfa (bireysel dosya).
    assert run.manifest["total_booklets"] == 8
    assert run.manifest["total_pages"] == 15

    pages = _room_pdf_pages(run)
    own = [text for text in pages if target.full_name in text]
    assert len(own) == 1, "öğrencinin tek sayfalık kendi kitapçığı bekleniyordu"
    assert BIREYSEL_IZ in own[0]  # içerik bireysel PDF'ten
    assert "Coğrafya" in own[0]  # bant ders adı öteki kitapçıklarla AYNI
    assert "Sayfa 1 / 1" in own[0]
    # Bireysel içerik YALNIZ o öğrencinin sayfasında.
    assert sum(1 for text in pages if BIREYSEL_IZ in text) == 1
    # Sıra korunur: kitapçıklar koltuk sırasında, bireysel kitapçık kendi sırasında.
    order = [
        a.full_name for a in SeatAssignment.objects.filter(session=session).order_by("seat_no")
    ]
    seen = [name for text in pages for name in order if name in text]
    assert list(dict.fromkeys(seen)) == order


def test_bantta_ogrenciyi_ayiran_isaret_yok(monkeypatch: pytest.MonkeyPatch) -> None:
    """Bireysel kitapçığın bandı, aynı dersin öteki kitapçığıyla BİREBİR aynı bağlamdan basılır.

    Bant `booklet._overlay_pages_context` sözlüğünden üretilir; iki öğrencinin ilk
    sayfa sözlüğü ad ve okul no DIŞINDA aynı olmalıdır (ders adı dahil).
    """
    session = _session()
    _select_with_file(session, "101", pages=2)  # sayfa sayısı da aynı → salt bant ölçülür
    captured: dict[str, Any] = {}
    original = booklet.build_room_package

    def _capture(room_name: str, specs: Any, docs: Any, info: Any, **kwargs: Any) -> Any:
        captured["specs"], captured["docs"] = list(specs), dict(docs)
        return original(room_name, specs, docs, info, **kwargs)

    monkeypatch.setattr(booklet, "build_room_package", _capture)
    run = services.request_booklet_run(session)
    assert run.status == BookletRunStatus.COMPLETED, run.error_message

    docs = captured["docs"]
    page_counts = {key: 2 for key in docs}
    context = booklet._overlay_pages_context(captured["specs"], docs, page_counts)

    def _first_page(number: str) -> dict[str, Any]:
        page = next(p for p in context if p["student_number"] == number and p["page_no"] == 1)
        return {k: v for k, v in page.items() if k not in {"full_name", "student_number"}}

    assert _first_page("101") == _first_page("102")
    # Bireysel doküman kendi anahtarıyla girer, isimsiz yedeğe kapalıdır; grup anahtarı DEĞİŞMEZ.
    individual = [doc for doc in docs.values() if not doc.backup]
    assert len(individual) == 1
    assert services_individual.INDIVIDUAL_KEY_SEP in individual[0].group_key
    assert not any(
        services_individual.INDIVIDUAL_KEY_SEP in group
        for group in SeatAssignment.objects.filter(session=session).values_list(
            "conflict_group", flat=True
        )
    )
    for text in _room_pdf_pages(run):
        assert "BEP" not in text and "ireysel" not in text


def test_secili_ama_dosyasiz_ogrenci_uretimi_durdurur() -> None:
    session = _session()
    student = _student("101")
    services_individual.add_iep_student(student.pk)
    services_individual.select_individual(session, student_id=student.pk)

    with pytest.raises(ValidationError) as exc_info:
        services.request_booklet_run(session)
    message = str(exc_info.value)
    assert "1 öğrencinin dosyası yüklenmemiş" in message
    # KVKK md. 6: hata metni öğrenci kimliği taşımaz (ad, okul no, pk).
    assert "AD0" not in message and "SOYAD" not in message and "101" not in message


def test_yalniz_bireysel_dosyali_grupta_ders_dosyasi_aranmaz() -> None:
    """Tek öğrencili grup (mazeret oturumu gibi): ders dosyası hiç yüklenmeyebilir."""
    sube(9, "A", students=1, start_no=101)
    sube(10, "A", students=3, start_no=201)
    session = oturum()
    services.add_session_course(
        session,
        course_id=ders("Coğrafya", levels=[9]).pk,
        participant_type=ParticipantType.LEVEL,
        level=9,
    )
    sc10 = services.add_session_course(
        session,
        course_id=ders("Fizik", levels=[10]).pk,
        participant_type=ParticipantType.LEVEL,
        level=10,
    )
    services.set_session_rooms(session, [{"room_id": salon("D-201", plan=PLAN_3X2_DOUBLE).pk}])
    services.distribute_session(session, seed=7)
    services.upload_question_document(sc10, file_bytes=_question_pdf(1, title="Fizik"))

    with pytest.raises(ValidationError, match="Soru dosyası eksik dersler: Coğrafya"):
        services.request_booklet_run(session)
    _select_with_file(session, "101")
    run = services.request_booklet_run(session)
    assert run.status == BookletRunStatus.COMPLETED, run.error_message
    assert run.manifest["total_booklets"] == 4


def test_isimsiz_yedek_bireysel_dosyadan_basilmaz() -> None:
    """Salona tek öğrenciye özgü sınavın ADSIZ kopyası düşmemeli (`CourseDoc.backup`)."""
    session = _session()
    _select_with_file(session, "101", pages=1)
    run = services.request_booklet_run(session, backup_copies=6)
    assert run.status == BookletRunStatus.COMPLETED, run.error_message
    pages = _room_pdf_pages(run)
    assert sum(1 for text in pages if BIREYSEL_IZ in text) == 1  # yalnız sahibinin kitapçığı
    assert run.manifest["total_booklets"] == 8 + 6


def test_course_doc_backup_varsayilani_eski_davranisi_korur() -> None:
    doc = booklet.CourseDoc(
        group_key="1:9",
        course_name="Coğrafya",
        pdf_bytes=b"",
        score_mode=ScoreMode.SINGLE_BOX,
        question_count=None,
    )
    assert doc.backup is True


def test_uretimden_sonra_degisen_bireysel_dosya_uretimi_bayatlatir() -> None:
    session = _session()
    run = services.request_booklet_run(session)
    assert BookletRunSerializer(run).data["is_stale"] is False

    row = _select_with_file(session, "101")
    assert BookletRunSerializer(run).data["is_stale"] is True

    fresh = services.request_booklet_run(session)
    assert BookletRunSerializer(fresh).data["is_stale"] is False
    # KALDIRMA da bayatlatır (satır KATI silinir; damga oturumda durur).
    services_individual.remove_individual(row)
    assert BookletRunSerializer(fresh).data["is_stale"] is True


def test_yeniden_dagitim_ve_taslaga_alma_secimi_korur() -> None:
    session = _session()
    _select_with_file(session, "101")
    session = services.revert_session_to_draft(session)
    assert session.status == ExamSessionStatus.DRAFT
    assert IndividualQuestionDocument.objects.filter(session=session).count() == 1

    services.distribute_session(session, seed=99)
    run = services.request_booklet_run(session)
    assert run.status == BookletRunStatus.COMPLETED, run.error_message
    assert run.manifest["individual_booklets"] == 1


# ===========================================================================
# İşaret yok — salonlara giden evrak
# ===========================================================================


def test_salon_evrakinda_ve_pakette_ayiran_isaret_yok() -> None:
    session = _session()
    _select_with_file(session, "101")

    for code in ("r1", "r4", "r7", "r8"):
        text = _pdf_text(services.render_session_report(session, code).content)
        assert "BEP" not in text and "ireysel" not in text, code

    workbook = load_workbook(io.BytesIO(services.render_session_report(session, "r5").content))
    cells = [
        str(cell.value)
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
    ]
    assert not any("BEP" in value or "ireysel" in value for value in cells)

    # İdare özeti "Tümünü indir" paketine GİRMEZ ve rapor kodu DEĞİLDİR.
    with zipfile.ZipFile(io.BytesIO(services.render_session_reports_zip(session).content)) as zf:
        assert not any("bep" in name.lower() for name in zf.namelist())
    assert not any("bep" in code or "iep" in code for code in services.REPORT_CODES)


def test_yerlesim_ve_ders_grubu_bireysel_dosyadan_etkilenmez() -> None:
    """Aynı seed → aynı dağıtım: bireysel dosya motor girdisi DEĞİLDİR."""
    session = _session()
    before = {
        a.student_id: (a.room_id, a.seat_no, a.conflict_group)
        for a in SeatAssignment.objects.filter(session=session)
    }
    _select_with_file(session, "101")
    session = services.revert_session_to_draft(session)
    services.distribute_session(session, seed=42)
    after = {
        a.student_id: (a.room_id, a.seat_no, a.conflict_group)
        for a in SeatAssignment.objects.filter(session=session)
    }
    assert after == before


# ===========================================================================
# İdare özeti
# ===========================================================================


def test_idare_ozeti_icerigi() -> None:
    session = _session()
    _select_with_file(session, "101")
    waiting = _student("201")
    services_individual.add_iep_student(waiting.pk)
    services_individual.select_individual(session, student_id=waiting.pk)
    services_individual.add_iep_student(_student("102").pk)  # listede, seçim yok

    report = services_individual.render_iep_summary(session)
    assert report.filename.endswith(".pdf") and report.content_type == "application/pdf"
    text = _pdf_text(report.content)
    assert "İdare Özeti" in text and "ÖZEL NİTELİKLİ" in text
    for number in ("101", "102", "201"):
        assert number in text
    assert "103" not in text  # listede olmayan öğrenci özete GİRMEZ
    assert "Bireysel soru dosyası" in text
    assert "Dersin soru dosyası" in text
    assert "dosya bekleniyor" in text
    # Dayanak depodaki metinlerden; md. 6/3 bendi BİLİNÇLE gösterilmez.
    assert "md. 4/1-ç" in text and "md. 5/1-u" in text and "md. 45/1-ğ" in text
    assert "6/3" not in text and "5/2" not in text


def test_idare_ozeti_kapilari() -> None:
    session = _session(with_course_docs=False)
    with pytest.raises(ValidationError, match="BEP kapsamında yerleşmiş öğrenci yok"):
        services_individual.render_iep_summary(session)
    with pytest.raises(ValidationError, match="Önce dağıtım"):
        services_individual.render_iep_summary(oturum(name="Taslak"))


def test_panel_satirlari_ve_yetim_satir() -> None:
    session = _session()
    _select_with_file(session, "101")
    services_individual.add_iep_student(_student("202").pk)

    rows = services_individual.session_rows(session)
    assert [row["seat_no"] for row in rows] == sorted(row["seat_no"] for row in rows)
    by_number = {row["student_number"]: row for row in rows}
    assert set(by_number) == {"101", "202"}
    assert by_number["101"]["document"]["has_file"] is True
    assert by_number["101"]["course_label"].startswith("Coğrafya")
    assert by_number["202"]["document"] is None
    # Satır özeti öğrenci kimliği taşımaz (opak satır kimliği).
    assert set(by_number["101"]["document"]) == {
        "id",
        "has_file",
        "page_count",
        "score_mode",
        "question_count",
    }

    # Öğrenci yerleşimden düşerse satır GÖRÜNÜR kalır (kaldırılabilsin) ve üretimi etkilemez.
    SeatAssignment.objects.filter(session=session, student=_student("101")).delete()
    orphan = next(
        r for r in services_individual.session_rows(session) if r["student_number"] == "101"
    )
    assert orphan["seat_no"] is None and orphan["room_name"] == ""
    assert services.request_booklet_run(session).manifest["individual_booklets"] == 0


# ===========================================================================
# KVKK — silme yolları
# ===========================================================================


def test_ayrilan_ogrencinin_kayitlari_kati_silinir(
    django_capture_on_commit_callbacks: Any,
) -> None:
    session = _session()
    row = _select_with_file(session, "101")
    storage, name = row.file.storage, row.file.name
    services.approve_session(session)  # onaylı oturum da KAPSAMDA (öğrenci ayrıldı)

    with django_capture_on_commit_callbacks(execute=True):
        persons.update_student(_student("101"), status=StudentStatus.LEFT)
    assert IepStudent.all_objects.count() == 0
    assert IndividualQuestionDocument.all_objects.count() == 0
    assert not storage.exists(name)


def test_silinen_ogrencinin_kayitlari_kati_silinir() -> None:
    session = _session()
    _select_with_file(session, "101")
    persons.delete_student(_student("101"))
    assert IepStudent.all_objects.count() == 0
    assert IndividualQuestionDocument.all_objects.count() == 0


def test_kancaya_ugramayan_durum_degisikligi_okumada_temizlenir() -> None:
    session = _session()
    _select_with_file(session, "101")
    services_individual.add_iep_student(_student("102").pk)
    Student.objects.filter(student_number="101").update(status=StudentStatus.LEFT)

    assert [i["student_number"] for i in services_individual.iep_list()] == ["102"]
    assert IndividualQuestionDocument.all_objects.count() == 0


def test_listeden_cikarma_yalniz_onaylanmamis_oturumu_temizler() -> None:
    approved = _session()
    _select_with_file(approved, "101")
    services.approve_session(approved)

    draft = oturum(name="İkinci Sınav")
    services.add_session_course(
        draft,
        course_id=ders("Coğrafya", levels=[9]).pk,
        participant_type=ParticipantType.LEVEL,
        level=9,
    )
    services.set_session_rooms(draft, [{"room_id": salon("D-202", plan=PLAN_3X2_DOUBLE).pk}])
    services.distribute_session(draft, seed=5)
    _select_with_file(draft, "101")

    row = services_individual.iep_rows()[_student("101").pk]
    services_individual.remove_iep_student(row)
    assert IndividualQuestionDocument.objects.filter(session=draft).count() == 0
    # Onaylı oturumun kaydı o sınavın yapıldığı hâlin parçasıdır — yerinde kalır.
    assert IndividualQuestionDocument.objects.filter(session=approved).count() == 1


def test_tumunu_sil(django_capture_on_commit_callbacks: Any) -> None:
    session = _session()
    row = _select_with_file(session, "101")
    services_individual.add_iep_student(_student("102").pk)
    storage, name = row.file.storage, row.file.name

    with django_capture_on_commit_callbacks(execute=True):
        result = services_individual.delete_all()
    assert result == {"students": 2, "documents": 1}
    assert IepStudent.all_objects.count() == 0
    assert IndividualQuestionDocument.all_objects.count() == 0
    assert not storage.exists(name)


def test_arsiv_anonimlestirme_ve_oturum_silme_satirlari_kaldirir(
    django_capture_on_commit_callbacks: Any,
) -> None:
    session = _session()
    row = _select_with_file(session, "101")
    storage, name = row.file.storage, row.file.name
    services.approve_session(session)
    services.archive_session(session)
    session.refresh_from_db()
    with django_capture_on_commit_callbacks(execute=True):
        counts = services.anonymize_exam_session(session)
    assert counts["individual_questions"] == 1
    assert IndividualQuestionDocument.all_objects.count() == 0
    assert not storage.exists(name)
    # BEP listesi oturumdan bağımsızdır; anonimleştirme ona dokunmaz.
    assert IepStudent.objects.count() == 1

    draft = _draft_with_individual_row()
    services.remove_exam_session(draft)
    assert IndividualQuestionDocument.all_objects.filter(session=draft).count() == 0


def _draft_with_individual_row() -> ExamSession:
    draft = oturum(name="Silinecek Taslak")
    IndividualQuestionDocument.objects.create(session=draft, student_ref="1")
    return draft


def test_gunluk_ogrenci_kimligi_tasimaz(caplog: pytest.LogCaptureFixture) -> None:
    session = _session()
    _select_with_file(session, "101")
    student = _student("101")
    with caplog.at_level(logging.DEBUG, logger="kelebek_sinav"):
        Student.objects.filter(pk=student.pk).update(status=StudentStatus.LEFT)
        services_individual.purge_stale()
        services_individual.delete_all()
    messages = [record.getMessage() for record in caplog.records]
    assert any("BEP kayıtları" in message for message in messages), "günlük yakalanmadı"
    for message in messages:
        assert "AD0" not in message and "SOYAD" not in message
        assert "101" not in message


# ===========================================================================
# API
# ===========================================================================


def test_api_akisi_uctan_uca() -> None:
    session = _session()
    student = _student("101")
    client = APIClient()

    assert client.get("/api/v1/iep-students/").data == {"results": []}
    created = client.post("/api/v1/iep-students/", {"student_id": student.pk}, format="json")
    assert created.status_code == 201
    duplicate = client.post("/api/v1/iep-students/", {"student_id": student.pk}, format="json")
    assert duplicate.status_code == 400
    listing = client.get("/api/v1/iep-students/").data["results"]
    assert [item["student_number"] for item in listing] == ["101"]

    rows = client.get(f"/api/v1/individual-questions/?session={session.pk}").data["rows"]
    assert len(rows) == 1 and rows[0]["document"] is None

    selected = client.post(
        "/api/v1/individual-questions/",
        {"session_id": session.pk, "student_id": student.pk},
        format="json",
    )
    assert selected.status_code == 201 and selected.data["has_file"] is False
    document_id = selected.data["id"]

    missing = client.get(f"/api/v1/individual-questions/{document_id}/file/")
    assert missing.status_code == 404

    uploaded = client.post(
        f"/api/v1/individual-questions/{document_id}/file/",
        {"file": io.BytesIO(_question_pdf(1, title=BIREYSEL_IZ)), "score_mode": "SINGLE_BOX"},
        format="multipart",
    )
    assert uploaded.status_code == 201, uploaded.data
    assert uploaded.data["has_file"] is True and uploaded.data["page_count"] == 1

    downloaded = client.get(f"/api/v1/individual-questions/{document_id}/file/")
    assert downloaded.status_code == 200
    # İndirilen dosyanın adı öğrenciyi ANMAZ.
    assert "101" not in downloaded["Content-Disposition"]

    summary = client.get(f"/api/v1/individual-questions/summary/?session={session.pk}")
    assert summary.status_code == 200 and summary["Content-Type"] == "application/pdf"

    booklets = client.post(f"/api/v1/exam-sessions/{session.pk}/booklets/", {}, format="json")
    assert booklets.status_code == 201
    assert booklets.data["manifest"]["individual_booklets"] == 1

    removed = client.delete(f"/api/v1/individual-questions/{document_id}/")
    assert removed.status_code == 204
    assert client.delete(f"/api/v1/individual-questions/{document_id}/").status_code == 404

    assert client.delete(f"/api/v1/iep-students/{listing[0]['id']}/").status_code == 204
    assert client.get("/api/v1/iep-students/").data == {"results": []}


def test_api_hatalari_ve_tumunu_sil() -> None:
    session = _session()
    client = APIClient()
    assert client.get("/api/v1/individual-questions/?session=999999").status_code == 404
    assert client.get("/api/v1/individual-questions/summary/").status_code == 404
    empty = client.get(f"/api/v1/individual-questions/summary/?session={session.pk}")
    assert empty.status_code == 400
    assert client.delete("/api/v1/iep-students/999999/").status_code == 404
    bad = client.post("/api/v1/iep-students/", {"student_id": 0}, format="json")
    assert bad.status_code == 400

    _select_with_file(session, "101")
    wiped = client.post("/api/v1/iep-students/delete-all/", {}, format="json")
    assert wiped.status_code == 200 and wiped.data == {"students": 1, "documents": 1}


def test_url_yollari_ogrenci_kimligi_tasimaz() -> None:
    """Django 4xx yanıtını YOLUYLA günlüğe yazar — öğrenci pk'si yolda olamaz.

    Yollardaki tek değişken SATIR kimliğidir (`pk`); öğrenci pk'si gövdede gelir.
    """
    from apps.sinav.urls import router

    patterns = [
        str(url.pattern)
        for url in router.urls
        if "iep-students" in str(url.pattern) or "individual-questions" in str(url.pattern)
    ]
    assert patterns
    for pattern in patterns:
        assert set(re.findall(r"[(][?]P<([a-z_]+)>", pattern)) <= {"pk", "format"}, pattern
