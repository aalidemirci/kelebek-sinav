"""Mazeret takibi ve mazeret sınavı testleri (19.09.2026, kullanıcı isteği).

Kurallar `services_makeup` modül başında; kısaca: mazeret sınavına YALNIZ
"Mazeretli" kayıt girer, bir defaya mahsustur (mazeret sınavında da girmeyene
ikincisi açılmaz), dönem içinde açılır, 5 iş günü süresi UYARIDIR. Hata ve uyarı
metinleri öğrenci ADI taşımaz (okul no).
"""

from __future__ import annotations

import io
from datetime import date, time

import pytest
from django.core.exceptions import ValidationError
from django.utils import timezone
from openpyxl import load_workbook
from pypdf import PdfReader
from rest_framework.test import APIClient

from apps.dersler.models import Course
from apps.okul.models import SchoolConfig, Student, StudentStatus
from apps.sinav import participants, services, services_makeup
from apps.sinav.models import (
    ExamAttendanceRecord,
    ExamRoom,
    ExamSession,
    ExamSessionCourse,
    ExamSessionStatus,
    ExamSessionType,
    ExcuseStatus,
    ParticipantType,
    SeatAssignment,
)
from apps.sinav.services_calendar import _validate_entry_participants
from apps.sinav.tests.oturum_yardim import dagitilmis_oturum, ders, donem, oturum, salon, sube

pytestmark = pytest.mark.django_db

URL = "/api/v1/makeup/"
MAZERET_TARIHI = date(2026, 11, 23)


# ---------------------------------------------------------------------------
# Kurulum
# ---------------------------------------------------------------------------
def _okul() -> tuple[Course, Course, ExamRoom]:
    """9/A (101-104) ve 10/B (201-204) — iki ders, 8 koltuklu tek salon."""
    sube(9, "A", students=4, start_no=101)
    sube(10, "B", students=4, start_no=201)
    return ders("Coğrafya", levels=[9]), ders("Fizik", levels=[10]), salon("D-201")


def _onayli(c9: Course, c10: Course, room: ExamRoom, **kwargs: object) -> ExamSession:
    """İki dersli, dağıtılmış ve ONAYLI oturum (yoklama yalnız onaylıda açılır)."""
    session = oturum(**kwargs)
    for course, level in ((c9, 9), (c10, 10)):
        services.add_session_course(
            session, course_id=course.pk, participant_type=ParticipantType.LEVEL, level=level
        )
    services.set_session_rooms(session, [{"room_id": room.pk}])
    session, _sonuc, rapor = services.distribute_session(session, seed=42)
    assert rapor.is_valid
    return services.approve_session(session)


def _girmedi(
    session: ExamSession, numara: str, durum: str = ExcuseStatus.PENDING, note: str = ""
) -> ExamAttendanceRecord:
    koltuk = next(
        s for s in SeatAssignment.objects.filter(session=session) if s.student_number == numara
    )
    return services.mark_absent(
        session, seat_assignment_id=koltuk.pk, excuse_status=durum, note=note
    )


def _mazeret_sinavi(*kayitlar: ExamAttendanceRecord, **kwargs: object) -> ExamSession:
    return services_makeup.create_makeup_session(
        record_ids=[k.pk for k in kayitlar],
        exam_date=kwargs.pop("exam_date", MAZERET_TARIHI),  # type: ignore[arg-type]
        start_time=time(10, 0),
        **kwargs,  # type: ignore[arg-type]
    )


def _dagit_onayla(session: ExamSession, room: ExamRoom) -> ExamSession:
    services.set_session_rooms(session, [{"room_id": room.pk}])
    session, _sonuc, rapor = services.distribute_session(session, seed=7)
    assert rapor.is_valid
    return services.approve_session(session)


def _satir(record: ExamAttendanceRecord) -> dict[str, object]:
    satirlar = services_makeup.absence_rows(donem().pk)
    return next(s for s in satirlar if s["record_id"] == record.pk)


# ---------------------------------------------------------------------------
# 5 iş günü (Yönerge md. 5/1-y) — UYARI, engel değil
# ---------------------------------------------------------------------------
@pytest.mark.parametrize(
    ("sinav", "son_gun"),
    [
        (date(2026, 11, 12), date(2026, 11, 19)),  # Perşembe → ertesi Perşembe
        (date(2026, 11, 13), date(2026, 11, 20)),  # Cuma → ertesi Cuma
        (date(2026, 11, 14), date(2026, 11, 20)),  # Cumartesi → sayım Pazartesi başlar
        (date(2026, 11, 16), date(2026, 11, 23)),  # Pazartesi → ertesi Pazartesi
    ],
)
def test_bildirim_son_gunu_hafta_sonunu_atlar(sinav: date, son_gun: date) -> None:
    assert services_makeup.notice_deadline(sinav) == son_gun


def test_suresi_gecen_yalniz_karar_bekleyende_isaretlenir(monkeypatch: pytest.MonkeyPatch) -> None:
    c9, c10, room = _okul()
    session = _onayli(c9, c10, room)  # sınav 16.11.2026 → son gün 23.11.2026
    bekleyen = _girmedi(session, "101")
    kararli = _girmedi(session, "102", ExcuseStatus.EXCUSED)
    monkeypatch.setattr(timezone, "localdate", lambda: date(2026, 11, 24))
    assert _satir(bekleyen)["notice_overdue"] is True
    assert _satir(kararli)["notice_overdue"] is False
    monkeypatch.setattr(timezone, "localdate", lambda: date(2026, 11, 23))
    assert _satir(bekleyen)["notice_overdue"] is False  # son gün dahil


# ---------------------------------------------------------------------------
# Takip listesi
# ---------------------------------------------------------------------------
def test_takip_satiri_ders_duzey_ve_durum_tasir() -> None:
    c9, c10, room = _okul()
    session = _onayli(c9, c10, room)
    kayit = _girmedi(session, "201", note="Rapor no 12, 16.11.2026")
    satir = _satir(kayit)
    assert satir["course_label"] == "Fizik — 10. Sınıf"
    assert satir["class_label"] == "10/B"
    assert satir["excuse_label"] == "Beklemede"
    assert satir["notice_deadline"] == "2026-11-23"
    assert satir["can_makeup"] is False  # karar verilmeden mazeret sınavına alınmaz
    services.update_attendance_record(kayit, excuse_status=ExcuseStatus.EXCUSED)
    assert _satir(kayit)["can_makeup"] is True

    ozet = services_makeup.absence_summary(services_makeup.absence_rows(donem().pk))
    assert ozet["total"] == 1 and ozet["excused"] == 1 and ozet["awaiting_makeup"] == 1


def test_silinen_oturumun_kaydi_listeye_girmez() -> None:
    c9, c10, room = _okul()
    session = _onayli(c9, c10, room)
    _girmedi(session, "101")
    ExamSession.objects.filter(pk=session.pk).update(deleted_at=session.updated_at)
    assert services_makeup.absence_rows(donem().pk) == []


# ---------------------------------------------------------------------------
# Mazeret sınavı oluşturma
# ---------------------------------------------------------------------------
def test_mazeret_sinavi_yalniz_mazeretli_ogrencilerle_kurulur() -> None:
    c9, c10, room = _okul()
    kaynak = _onayli(c9, c10, room)
    a = _girmedi(kaynak, "101", ExcuseStatus.EXCUSED)
    b = _girmedi(kaynak, "201", ExcuseStatus.EXCUSED)
    _girmedi(kaynak, "102", ExcuseStatus.UNEXCUSED)

    mazeret = _mazeret_sinavi(a, b)
    assert mazeret.is_makeup and mazeret.status == ExamSessionStatus.DRAFT
    assert mazeret.name == services_makeup.DEFAULT_MAKEUP_NAME
    assert mazeret.semester_id == kaynak.semester_id
    satirlar = {
        (sc.course_id, sc.level): sc for sc in ExamSessionCourse.objects.filter(session=mazeret)
    }
    assert set(satirlar) == {(c9.pk, 9), (c10.pk, 10)}
    assert all(sc.participant_type == ParticipantType.MAKEUP for sc in satirlar.values())

    cozum = participants.resolve_session(mazeret)
    assert sorted(p.student_number for p in cozum.participants) == ["101", "201"]
    assert {p.conflict_group for p in cozum.participants} == {f"{c9.pk}:9", f"{c10.pk}:10"}
    assert _satir(a)["makeup_session_id"] == mazeret.pk and _satir(a)["can_makeup"] is False


def test_mazereti_kabul_edilmeyen_reddedilir_ad_yazilmaz() -> None:
    c9, c10, room = _okul()
    kaynak = _onayli(c9, c10, room)
    bekleyen = _girmedi(kaynak, "103")
    with pytest.raises(ValidationError) as exc:
        _mazeret_sinavi(bekleyen)
    metin = str(exc.value)
    assert "Okul No 103" in metin and "Mazeretli" in metin
    assert "AD2" not in metin and "SOYAD" not in metin  # KVKK: ad değil okul no
    assert not ExamSession.objects.filter(is_makeup=True).exists()


def test_zaten_alinan_ve_iki_sinavi_birden_secilen_reddedilir() -> None:
    c9, c10, room = _okul()
    ilk = _onayli(c9, c10, room)
    ikinci = _onayli(c9, c10, room, name="2. Ortak Sınav", exam_date=date(2026, 11, 17))
    a = _girmedi(ilk, "101", ExcuseStatus.EXCUSED)
    a2 = _girmedi(ikinci, "101", ExcuseStatus.EXCUSED)
    with pytest.raises(ValidationError, match="iki sınavın mazeretine birden"):
        _mazeret_sinavi(a, a2)
    _mazeret_sinavi(a)
    with pytest.raises(ValidationError, match="zaten 'Mazeret Sınavı'"):
        _mazeret_sinavi(a)
    # Farklı günlerin sınavları AYRI mazeret oturumlarında (kullanıcı kararı: aynı oturuma
    # farklı günlerin dersleri de girebilir — ama aynı öğrenci iki derse giremez).
    assert _mazeret_sinavi(a2, name="Mazeret Sınavı 2").is_makeup


def test_farkli_donemlerin_kayitlari_tek_mazeret_sinavina_girmez() -> None:
    c9, c10, room = _okul()
    birinci = _onayli(c9, c10, room)
    ikinci = _onayli(
        c9,
        c10,
        room,
        name="2. Dönem 1. Sınav",
        exam_date=date(2027, 3, 10),
        term_id=donem(sequence=2).pk,
    )
    a = _girmedi(birinci, "101", ExcuseStatus.EXCUSED)
    b = _girmedi(ikinci, "202", ExcuseStatus.EXCUSED)
    with pytest.raises(ValidationError, match="farklı dönemlerin"):
        _mazeret_sinavi(a, b)


def test_mazeret_sinavinda_da_girmeyene_ikincisi_acilmaz(monkeypatch: pytest.MonkeyPatch) -> None:
    """OKY md. 48/1 "bir defaya mahsus"; sonuç sütunu kaynak satırda "Girmedi"."""
    c9, c10, room = _okul()
    kaynak = _onayli(c9, c10, room)
    a = _girmedi(kaynak, "101", ExcuseStatus.EXCUSED)
    b = _girmedi(kaynak, "102", ExcuseStatus.EXCUSED)
    mazeret = _dagit_onayla(_mazeret_sinavi(a, b), room)
    assert _satir(a)["makeup_result"] == "pending"  # sınav günü gelmedi

    tekrar = _girmedi(mazeret, "101", ExcuseStatus.EXCUSED)
    with pytest.raises(ValidationError, match="bir defaya mahsus"):
        _mazeret_sinavi(tekrar)
    assert _satir(tekrar)["session_is_makeup"] is True
    assert _satir(tekrar)["can_makeup"] is False
    assert _satir(a)["makeup_result_label"] == "Girmedi"

    monkeypatch.setattr(timezone, "localdate", lambda: MAZERET_TARIHI)
    assert _satir(b)["makeup_result"] == "attended"


def test_mazeretsiz_cekilen_duser_ve_yeniden_dagitim_istenir() -> None:
    c9, c10, room = _okul()
    kaynak = _onayli(c9, c10, room)
    a = _girmedi(kaynak, "101", ExcuseStatus.EXCUSED)
    b = _girmedi(kaynak, "102", ExcuseStatus.EXCUSED)
    mazeret = _mazeret_sinavi(a, b)
    services.set_session_rooms(mazeret, [{"room_id": room.pk}])
    mazeret, _sonuc, _rapor = services.distribute_session(mazeret, seed=7)

    services.update_attendance_record(b, excuse_status=ExcuseStatus.UNEXCUSED)
    cozum = participants.resolve_session(mazeret)
    assert [p.student_number for p in cozum.participants] == ["101"]
    uyarilar = " ".join(w for c in cozum.courses for w in c.warnings)
    assert "1 öğrencinin mazeret durumu artık 'Mazeretli' değil" in uyarilar
    kayma = participants.placement_drift(mazeret, cozum)
    assert kayma is not None and kayma.removed == 1


def test_ayrilan_ogrenci_mazeret_sinavindan_duser() -> None:
    c9, c10, room = _okul()
    kaynak = _onayli(c9, c10, room)
    a = _girmedi(kaynak, "104", ExcuseStatus.EXCUSED)
    mazeret = _mazeret_sinavi(a)
    Student.objects.filter(pk=a.student_id).update(status=StudentStatus.LEFT)
    cozum = participants.resolve_session(mazeret)
    assert cozum.participants == []
    assert any("artık aktif değil" in w for c in cozum.courses for w in c.warnings)


def test_mazeret_sinavindan_cikarma_ve_silinen_oturum() -> None:
    c9, c10, room = _okul()
    kaynak = _onayli(c9, c10, room)
    a = _girmedi(kaynak, "101", ExcuseStatus.EXCUSED)
    mazeret = _mazeret_sinavi(a)

    assert services_makeup.remove_from_makeup(record_ids=[a.pk]) == 1
    assert _satir(a)["can_makeup"] is True and _satir(a)["makeup_session_id"] is None
    with pytest.raises(ValidationError, match="mazeret sınavına alınmamış"):
        services_makeup.remove_from_makeup(record_ids=[a.pk])

    # Taslak mazeret oturumu silinince kayıt kendiliğinden serbest kalır.
    ikinci = _mazeret_sinavi(a, name="Mazeret Sınavı 2")
    services.remove_exam_session(ikinci)
    assert _satir(a)["can_makeup"] is True
    assert mazeret.pk != ikinci.pk


def test_yapilmis_mazeret_sinavindan_cikarilamaz() -> None:
    c9, c10, room = _okul()
    kaynak = _onayli(c9, c10, room)
    a = _girmedi(kaynak, "101", ExcuseStatus.EXCUSED)
    _dagit_onayla(_mazeret_sinavi(a), room)
    with pytest.raises(ValidationError, match="onaylanmış"):
        services_makeup.remove_from_makeup(record_ids=[a.pk])


# ---------------------------------------------------------------------------
# Kapılar: mazeret satırı elle kurulmaz; takvimde üçüncü tip yok; dönem sabit
# ---------------------------------------------------------------------------
def test_mazeret_satiri_elle_eklenmez_ve_yalniz_suresi_degisir() -> None:
    c9, c10, room = _okul()
    kaynak = _onayli(c9, c10, room)
    mazeret = _mazeret_sinavi(_girmedi(kaynak, "101", ExcuseStatus.EXCUSED))
    with pytest.raises(ValidationError, match="Mazeret Takibi ekranından"):
        services.add_session_course(
            mazeret, course_id=c10.pk, participant_type=ParticipantType.LEVEL, level=10
        )
    normal = oturum(name="Deneme", exam_date=date(2026, 12, 1))
    with pytest.raises(ValidationError, match="Mazeret Takibi ekranından"):
        services.add_session_course(
            normal, course_id=c9.pk, participant_type=ParticipantType.MAKEUP, level=9
        )
    satir = ExamSessionCourse.objects.get(session=mazeret)
    with pytest.raises(ValidationError, match="Mazeret Takibi ekranından"):
        services.update_session_course(satir, level=10)
    # Dersin oturum içi nitelikleri (süre, "aynı kitapçık") mazeret satırında da açık.
    services.update_session_course(satir, duration_minutes=60, shared_booklet=True)
    satir.refresh_from_db()
    assert satir.duration_minutes == 60 and satir.shared_booklet is True
    assert (
        _satir(ExamAttendanceRecord.objects.get(makeup_course=satir))["makeup_session_status"]
        == ExamSessionStatus.DRAFT
    )

    # Plan kopyalama mazeret satırını gerekçesiyle ATLAR (patlamaz).
    rapor = services.copy_session_plan(normal, source_id=mazeret.pk, rooms=False)
    assert rapor["courses_created"] == []
    assert "Mazeret Takibi" in rapor["courses_skipped"][0]


def test_takvimde_mazeretli_kapsam_yok() -> None:
    with pytest.raises(ValidationError, match="yalnız mazeret sınavı oturumunda"):
        _validate_entry_participants(level=9, participant_type="MAKEUP", section_ids=None)


def test_mazeret_sinavinin_donemi_degismez() -> None:
    c9, c10, room = _okul()
    kaynak = _onayli(c9, c10, room)
    mazeret = _mazeret_sinavi(_girmedi(kaynak, "101", ExcuseStatus.EXCUSED))
    with pytest.raises(ValidationError, match="dönemi değiştirilemez"):
        services.update_exam_session(mazeret, term_id=donem(sequence=2).pk)
    # Aynı dönemin geri gönderilmesi (form) serbest.
    services.update_exam_session(mazeret, term_id=kaynak.semester_id, name="Mazeret — Kasım")


# ---------------------------------------------------------------------------
# Rapor (PDF + Excel) ve API
# ---------------------------------------------------------------------------
def _rapor_senaryosu() -> tuple[ExamAttendanceRecord, ...]:
    SchoolConfig.objects.create(
        pk=SchoolConfig.SINGLETON_PK,
        school_name="Örnek Anadolu Lisesi",
        principal_name="Şükrü Ağaoğlu",
    )
    c9, c10, room = _okul()
    okul_sinavi = _onayli(c9, c10, room)
    ulke = _onayli(
        c9,
        c10,
        room,
        name="Ülke Geneli Sınav",
        exam_date=date(2026, 11, 11),
        session_type=ExamSessionType.NATIONAL,
    )
    g = _girmedi(okul_sinavi, "101", ExcuseStatus.UNEXCUSED)
    bekleyen = _girmedi(okul_sinavi, "201", ExcuseStatus.EXCUSED, note="Hastane raporu 14.11.2026")
    ulusal = _girmedi(ulke, "202", ExcuseStatus.EXCUSED)
    return g, bekleyen, ulusal


def test_rapor_pdf_dort_bolum_ve_imza() -> None:
    g, bekleyen, ulusal = _rapor_senaryosu()
    pdf = services_makeup.render_makeup_report_pdf(donem().pk)
    reader = PdfReader(io.BytesIO(pdf))
    assert float(reader.pages[0].mediabox.width) > float(reader.pages[0].mediabox.height)
    metin = " ".join(" ".join(p.extract_text() or "" for p in reader.pages).split())
    for baslik in (
        "SINAVA GİRMEYEN ÖĞRENCİLER VE MAZERET TAKİP ÇİZELGESİ",
        "A. SINAVA GİRMEYEN ÖĞRENCİLER",
        "B. E-OKUL'A “G” İŞLENECEKLER",
        "C. MAZERET SINAVI BEKLEYENLER",
        "D. İL/İLÇE MİLLÎ EĞİTİM MÜDÜRLÜĞÜNE BİLDİRİLECEKLER",
        "Düzenleyen — Müdür Yardımcısı",
        "Şükrü Ağaoğlu",
        "Örnek Anadolu Lisesi Müdürlüğü",
        "6698 sayılı Kanunun 5/2-ç",
    ):
        assert baslik in metin, baslik
    assert "Hastane raporu 14.11.2026" in metin
    assert "Ülke geneli" in metin


def test_rapor_excel_bolumleri_ayri_sayfalarda() -> None:
    g, bekleyen, ulusal = _rapor_senaryosu()
    wb = load_workbook(io.BytesIO(services_makeup.build_makeup_workbook(donem().pk)))
    assert wb.sheetnames == ["Girmeyenler", "e-Okul G", "Mazeret bekleyenler", "İl-İlçe bildirimi"]

    def okul_nolari(sayfa: str, sutun: int) -> list[str]:
        ws = wb[sayfa]
        return [str(ws.cell(row=r, column=sutun).value) for r in range(5, ws.max_row + 1)]

    assert sorted(okul_nolari("Girmeyenler", 4)) == ["101", "201", "202"]
    assert okul_nolari("e-Okul G", 1) == ["101"]
    assert sorted(okul_nolari("Mazeret bekleyenler", 1)) == ["201", "202"]
    assert okul_nolari("İl-İlçe bildirimi", 1) == ["202"]
    # Sınav türü sütunu PDF'teki ifadeyle aynı ("Ülke" değil, "Ülke geneli").
    assert okul_nolari("İl-İlçe bildirimi", 5) == ["Ülke geneli"]


def test_api_liste_olustur_cikar_ve_rapor() -> None:
    g, bekleyen, ulusal = _rapor_senaryosu()
    client = APIClient()

    liste = client.get(f"{URL}absences/")
    assert liste.status_code == 200
    assert liste.data["semester_id"] == donem().pk
    assert liste.data["summary"]["total"] == 3
    assert liste.data["notice_business_days"] == 5

    ret = client.post(
        f"{URL}sessions/",
        {"record_ids": [g.pk], "exam_date": "2026-11-23", "start_time": "10:00"},
        format="json",
    )
    assert ret.status_code == 400 and "Mazeretli" in str(ret.data)

    yanit = client.post(
        f"{URL}sessions/",
        {
            "record_ids": [bekleyen.pk, ulusal.pk],
            "name": "Kasım Mazeret Sınavı",
            "exam_date": "2026-11-23",
            "start_time": "10:00",
            "duration_minutes": 40,
        },
        format="json",
    )
    assert yanit.status_code == 201, yanit.data
    assert yanit.data["is_makeup"] is True
    assert {c["participant_type"] for c in yanit.data["courses"]} == {"MAKEUP"}

    cikar = client.post(f"{URL}remove/", {"record_ids": [ulusal.pk]}, format="json")
    assert cikar.status_code == 200 and cikar.data == {"removed": 1}

    pdf = client.get(f"{URL}report/?kind=pdf")
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    assert f'filename="mazeret_takip_donem_{donem().pk}.pdf"' in pdf["Content-Disposition"]
    xlsx = client.get(f"{URL}report/?kind=xlsx&semester={donem().pk}")
    assert xlsx.status_code == 200 and xlsx.content.startswith(b"PK")
    assert client.get(f"{URL}report/?kind=doc").status_code == 400


def test_api_donem_yokken_bos_liste() -> None:
    yanit = APIClient().get(f"{URL}absences/")
    assert yanit.status_code == 200
    assert yanit.data["semester_id"] is None and yanit.data["rows"] == []


def test_ortak_kitapcikli_oturumda_duzey_sube_etiketinden() -> None:
    """`<ders>:*` grubunda (tüm seviyeler aynı kitapçık) düzey şube etiketinden çözülür."""
    session = dagitilmis_oturum()
    course = ders("Coğrafya")
    ExamSessionCourse.objects.filter(session=session).update(shared_booklet=True)
    SeatAssignment.objects.filter(session=session).update(conflict_group=f"{course.pk}:*")
    # Onay doğrulayıcısı tek gruba çevrilmiş yerleşimi reddederdi; yoklama için durum yeter.
    ExamSession.objects.filter(pk=session.pk).update(status=ExamSessionStatus.APPROVED)
    session.refresh_from_db()
    kayit = _girmedi(session, "201", ExcuseStatus.EXCUSED)
    assert _satir(kayit)["course_label"] == "Coğrafya — 10. Sınıf"
    assert _satir(kayit)["can_makeup"] is True
