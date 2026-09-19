"""Mazeret takibi ve mazeret sınavı (19.09.2026, kullanıcı isteği).

Kaynak, sınava girmeyen öğrenci kayıtlarıdır (`ExamAttendanceRecord`): her
oturumun Yoklama sekmesinde işaretlenir; mazeret durumu (Beklemede / Mazeretli /
Mazeretsiz) ve belge notu orada ya da Mazeret Takibi ekranında girilir. Bu modül
dönemin bütün kayıtlarını tek listede toplar, mazeret sınavı oturumu açar ve
takip raporunun bağlamını kurar. Kişisel veri YAZMAZ; yalnız okur ve bağlar.

Kurallar (metinler `docs/mevzuat/`; kararlar kullanıcıdan, 19.09.2026).
OKY = Ortaöğretim Kurumları Yönetmeliği, Yönerge = Yazılı ve Uygulamalı
Sınavlar Yönergesi:

- Mazeret sınavı özrünü belgelendiren öğrenci içindir (OKY md. 48/1, Yönerge
  md. 5/1-aa) → oturuma YALNIZ "Mazeretli" kayıtlar alınır; katılımcılar
  kayıttan anlık türetilir (`participants._resolve_makeup`), sonradan
  "Mazeretsiz"e çekilen düşer. Süre dönemi aşamaz (OKY md. 48/1) → mazeret
  sınavı kaydın döneminde açılır.
- Bildirim sınav tarihinden itibaren 5 iş günü içindedir (Yönerge md. 5/1-y;
  OKY md. 36/7 zorunlu hâlde okul yönetimine 20 iş gününe kadar uzatma
  yetkisi verir). Süre UYARIDIR, engel değil — karar okul müdürlüğünündür.
  Programda resmî tatil verisi yok: süre hafta sonları düşülerek hesaplanır
  (bayram arada kalırsa gerçek süre daha uzundur).
- Mazeret sınavı "bir defaya mahsus" yapılır (OKY md. 48/1); ülke/il/ilçe
  geneli sınavlar için ayrıca Yönerge md. 5/1-çç → mazeret sınavına da
  girmeyene İKİNCİ mazeret sınavı açılmaz.
- Geçerli mazereti olmadan sınava ya da mazeret sınavına girmeyenin puan
  hanesine "G" yazılır (OKY md. 48/4, Yönerge md. 6/1-f) → raporun "e-Okul'a
  G" bölümü: durumu "Mazeretsiz" olan her kayıt.
- Ülke / il / ilçe geneli sınava katılamayıp mazeret sınavına alınmasına karar
  verilenler il/ilçe millî eğitim müdürlüğüne resmî yazıyla bildirilir
  (Yönerge md. 5/1-z) → raporun bildirim bölümü (oturum türü "Okul" dışı).
  Ülke ve il geneli sınavların mazeret sınavı TARİHİ il MEM'ce ilan edilir
  (md. 5/1-aa) — oturum tarihi o ilana göre seçilir.
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, timedelta
from typing import TYPE_CHECKING, Any

from django.core.exceptions import ValidationError
from django.db import transaction
from django.utils import timezone

from apps.okul import selectors as okul_selectors
from apps.sinav.models import (
    ExamAttendanceRecord,
    ExamSession,
    ExamSessionCourse,
    ExamSessionStatus,
    ExamSessionType,
    ExcuseStatus,
    ParticipantType,
    SeatAssignment,
)

if TYPE_CHECKING:
    from apps.sinav.services import ReportFile

#: Veli bildirim süresi (Yönerge md. 5/1-y).
NOTICE_BUSINESS_DAYS = 5
DEFAULT_MAKEUP_NAME = "Mazeret Sınavı"


def notice_deadline(exam_date: date, days: int = NOTICE_BUSINESS_DAYS) -> date:
    """Sınav tarihinden sonraki `days`. iş günü (Pazartesi-Cuma).

    "Uygulama tarihinden itibaren en geç 5 iş günü içerisinde" — sayım sınavın
    ertesi iş gününden başlar: Perşembe sınavının son günü ertesi Perşembedir.
    """
    gun = exam_date
    kalan = days
    while kalan:
        gun += timedelta(days=1)
        if gun.weekday() < 5:
            kalan -= 1
    return gun


def live_makeup_course(record: ExamAttendanceRecord) -> ExamSessionCourse | None:
    """Kaydın bağlı olduğu mazeret satırı — satır ya da oturumu silinmişse None.

    İleri FK soft-delete'i SÜZMEZ ve soft-delete SET_NULL'ı tetiklemez
    (CLAUDE.md §2-§3): canlılık ELLE sorulur.
    """
    sc = record.makeup_course
    if sc is None or sc.deleted_at is not None or sc.session.deleted_at is not None:
        return None
    return sc


def _group_key(conflict_group: str, class_label: str) -> tuple[int, int] | None:
    """'<ders>:<düzey>' / '<ders>:*' → (ders, düzey). Joker düzey şube etiketinden."""
    ders, _, duzey = conflict_group.partition(":")
    try:
        course_id = int(ders)
    except ValueError:
        return None
    if duzey and duzey != "*":
        try:
            return course_id, int(duzey)
        except ValueError:
            return None
    bas = class_label.split("/", 1)[0].strip()
    return (course_id, int(bas)) if bas.isdigit() else None


def _seat_groups(records: list[ExamAttendanceRecord]) -> dict[tuple[int, int], tuple[int, int]]:
    """(oturum, öğrenci) → (ders, düzey) — aynı oturumdaki yerleşim kaydından (tek sorgu)."""
    oturumlar = {r.session_id for r in records}
    ogrenciler = {r.student_id for r in records if r.student_id is not None}
    sonuc: dict[tuple[int, int], tuple[int, int]] = {}
    for seat in SeatAssignment.objects.filter(
        session_id__in=oturumlar, student_id__in=ogrenciler
    ).only("session_id", "student_id", "conflict_group", "class_label"):
        anahtar = _group_key(seat.conflict_group, seat.class_label)
        if anahtar is not None and seat.student_id is not None:
            sonuc[(seat.session_id, seat.student_id)] = anahtar
    return sonuc


def _course_label(course_names: dict[int, str], course_id: int, level: int) -> str:
    from apps.dersler.text import level_label

    return f"{course_names.get(course_id, 'Ders')} — {level_label(level)}"


@dataclass(frozen=True)
class _MakeupInfo:
    session: ExamSession
    course: ExamSessionCourse


def _makeup_result(
    info: _MakeupInfo | None, student_id: int | None, absent_in: set[tuple[int, int]]
) -> str | None:
    """Mazeret sınavının sonucu: `absent` · `attended` · `pending` (henüz yapılmadı)."""
    if info is None or student_id is None:
        return None
    if (info.session.pk, student_id) in absent_in:
        return "absent"
    yapildi = info.session.status in (ExamSessionStatus.APPROVED, ExamSessionStatus.ARCHIVED)
    if yapildi and info.session.exam_date <= timezone.localdate():
        return "attended"
    return "pending"


_RESULT_LABELS = {
    "absent": "Girmedi",
    "attended": "Girdi",
    "pending": "Yapılmadı",
}


def absence_rows(semester_id: int) -> list[dict[str, Any]]:
    """Dönemin sınava girmeyen bütün kayıtları — takip ekranı ve rapor AYNI satırı kullanır.

    Anonimleştirilmiş arşiv kayıtları (öğrencisi boş) listeye girmez. Sıra: sınav
    tarihi, ders etiketi, okul no.
    """
    from apps.dersler import selectors as ders_selectors

    records = list(
        ExamAttendanceRecord.objects.filter(
            session__semester_id=semester_id,
            session__deleted_at__isnull=True,
            student__isnull=False,
        ).select_related("session", "room", "makeup_course", "makeup_course__session")
    )
    gruplar = _seat_groups(records)
    ders_adlari = ders_selectors.course_names_by_ids({cid for cid, _ in gruplar.values()})
    # Mazeret oturumlarında da girmeyen (öğrenci, oturum) çiftleri — sonuç sütunu.
    mazeret_oturumlari = {
        sc.session_id for r in records if (sc := live_makeup_course(r)) is not None
    }
    girmeyen_mazeret = set(
        ExamAttendanceRecord.objects.filter(session_id__in=mazeret_oturumlari).values_list(
            "session_id", "student_id"
        )
    )
    bugun = timezone.localdate()
    satirlar: list[dict[str, Any]] = []
    for r in records:
        oturum = r.session
        grup = gruplar.get((oturum.pk, r.student_id or 0))
        sc = live_makeup_course(r)
        info = _MakeupInfo(session=sc.session, course=sc) if sc is not None else None
        son_gun = notice_deadline(oturum.exam_date)
        sonuc = _makeup_result(info, r.student_id, girmeyen_mazeret)
        satirlar.append(
            {
                "record_id": r.pk,
                "session_id": oturum.pk,
                "session_name": oturum.name,
                "exam_date": oturum.exam_date.isoformat(),
                "session_is_makeup": oturum.is_makeup,
                "session_type": oturum.session_type,
                "session_type_label": oturum.get_session_type_display(),
                "external": oturum.session_type != ExamSessionType.SCHOOL,
                "course_id": grup[0] if grup else None,
                "level": grup[1] if grup else None,
                "course_label": (
                    _course_label(ders_adlari, grup[0], grup[1]) if grup else "(ders bulunamadı)"
                ),
                "student_id": r.student_id,
                "student_number": r.student_number,
                "full_name": r.full_name,
                "class_label": r.class_label,
                "excuse_status": r.excuse_status,
                "excuse_label": r.get_excuse_status_display(),
                "note": r.note,
                "notice_deadline": son_gun.isoformat(),
                # Süre yalnız KARAR BEKLEYEN kayıtta uyarıdır (karar verilmişse bilgi).
                "notice_overdue": r.excuse_status == ExcuseStatus.PENDING and bugun > son_gun,
                "makeup_session_id": info.session.pk if info else None,
                "makeup_session_name": info.session.name if info else "",
                "makeup_session_status": info.session.status if info else None,
                "makeup_date": info.session.exam_date.isoformat() if info else None,
                "makeup_result": sonuc,
                "makeup_result_label": _RESULT_LABELS.get(sonuc or "", ""),
                "can_makeup": (
                    not oturum.is_makeup
                    and r.excuse_status == ExcuseStatus.EXCUSED
                    and info is None
                    and grup is not None
                ),
            }
        )
    satirlar.sort(key=lambda s: (s["exam_date"], s["course_label"], s["student_number"]))
    return satirlar


def absence_summary(rows: list[dict[str, Any]]) -> dict[str, int]:
    """Takip ekranının sayaçları (rapor başlığıyla aynı)."""
    return {
        "total": len(rows),
        "pending": sum(1 for r in rows if r["excuse_status"] == ExcuseStatus.PENDING),
        "excused": sum(1 for r in rows if r["excuse_status"] == ExcuseStatus.EXCUSED),
        "unexcused": sum(1 for r in rows if r["excuse_status"] == ExcuseStatus.UNEXCUSED),
        "overdue": sum(1 for r in rows if r["notice_overdue"]),
        "awaiting_makeup": sum(1 for r in rows if r["can_makeup"]),
        "in_makeup": sum(1 for r in rows if r["makeup_session_id"] is not None),
    }


def semester_options() -> list[dict[str, Any]]:
    """Aktif ders yılının dönemleri + varsayılan (bugünü içeren, yoksa son dönem)."""
    yil = okul_selectors.active_school_year()
    if yil is None:
        return []
    bugun = timezone.localdate()
    donemler = list(okul_selectors.school_terms(school_year_id=yil.pk))
    varsayilan = next(
        (d.pk for d in donemler if d.start_date <= bugun <= d.end_date),
        donemler[-1].pk if donemler else None,
    )
    return [{"id": d.pk, "label": str(d), "default": d.pk == varsayilan} for d in donemler]


@transaction.atomic
def create_makeup_session(
    *,
    record_ids: list[int],
    name: str = DEFAULT_MAKEUP_NAME,
    exam_date: date,
    start_time: Any,
    duration_minutes: int = 40,
) -> ExamSession:
    """Seçilen "Mazeretli" kayıtlarla TASLAK mazeret sınavı oturumu açar.

    Her (ders, sınıf düzeyi) için bir "Mazeretli öğrenciler" satırı oluşur ve
    kayıt o satıra bağlanır; salon, dağıtım ve evrak normal oturum akışıdır
    (sihirbaz). Reddedilenler — hepsi okul numarasıyla söylenir, ad yazılmaz:
    mazereti kabul edilmemiş kayıt, zaten mazeret sınavına alınmış kayıt,
    mazeret sınavında da girmeyen (ikinci mazeret sınavı yok) ve başka dönemin
    kaydı (mazeret sınavı dönem içinde açılır).
    """
    from apps.sinav import services

    temiz = list(dict.fromkeys(int(i) for i in record_ids))
    if not temiz:
        raise ValidationError("Mazeret sınavına alınacak öğrenci seçin.")
    records = list(
        ExamAttendanceRecord.objects.filter(pk__in=temiz).select_related(
            "session", "makeup_course", "makeup_course__session"
        )
    )
    if len(records) != len(temiz):
        raise ValidationError("Seçilen kayıtlardan biri bulunamadı; listeyi yenileyin.")
    for r in records:
        etiket = f"Okul No {r.student_number}"
        if r.student_id is None:
            raise ValidationError(f"{etiket}: anonimleştirilmiş kayıt mazeret sınavına alınamaz.")
        if r.session.deleted_at is not None:
            raise ValidationError(f"{etiket}: girmediği sınavın oturumu silinmiş.")
        if r.session.is_makeup:
            raise ValidationError(
                f"{etiket}: mazeret sınavına da girmemiş — mazeret sınavı bir defaya mahsus "
                "yapılır (Ortaöğretim Kurumları Yönetmeliği md. 48/1)."
            )
        if r.excuse_status != ExcuseStatus.EXCUSED:
            raise ValidationError(
                f"{etiket}: mazeret durumu 'Mazeretli' değil — mazeret sınavına yalnız mazereti "
                "kabul edilen öğrenciler alınır."
            )
        mevcut = live_makeup_course(r)
        if mevcut is not None:
            raise ValidationError(
                f"{etiket}: zaten '{mevcut.session.name}' mazeret sınavına alınmış."
            )
    donemler = {r.session.semester_id for r in records}
    if len(donemler) > 1:
        raise ValidationError(
            "Seçilen öğrenciler farklı dönemlerin sınavlarından; mazeret sınavı dönem içinde açılır."
        )
    # Oturum tek zaman dilimidir: iki sınavın mazeretine birden seçilen öğrenci aynı
    # anda iki sınava giremez (dağıtım da bunu "iki derse düşüyor" diye durdururdu).
    gorulen: set[int] = set()
    for r in records:
        if r.student_id in gorulen:
            raise ValidationError(
                f"Okul No {r.student_number}: iki sınavın mazeretine birden seçildi; bir "
                "oturumda tek sınava girilir — bu dersleri ayrı mazeret sınavlarına alın."
            )
        gorulen.add(r.student_id or 0)
    gruplar = _seat_groups(records)
    satir_anahtari: dict[int, tuple[int, int]] = {}
    for r in records:
        grup = gruplar.get((r.session_id, r.student_id or 0))
        if grup is None:
            raise ValidationError(
                f"Okul No {r.student_number}: girmediği sınavın yerleşim kaydı bulunamadı."
            )
        satir_anahtari[r.pk] = grup

    oturum = services.create_exam_session(
        name=name or DEFAULT_MAKEUP_NAME,
        exam_date=exam_date,
        start_time=start_time,
        duration_minutes=duration_minutes,
        term_id=donemler.pop(),
    )
    oturum.is_makeup = True
    oturum.save(update_fields=["is_makeup", "updated_at"])
    satirlar: dict[tuple[int, int], ExamSessionCourse] = {}
    for grup in sorted(set(satir_anahtari.values())):
        satirlar[grup] = ExamSessionCourse.objects.create(
            session=oturum,
            course_id=grup[0],
            level=grup[1],
            participant_type=ParticipantType.MAKEUP,
            section_ids=[],
        )
    for r in records:
        r.makeup_course = satirlar[satir_anahtari[r.pk]]
        r.save(update_fields=["makeup_course", "updated_at"])
    return oturum


@transaction.atomic
def remove_from_makeup(*, record_ids: list[int]) -> int:
    """Seçilen kayıtları mazeret sınavından çıkarır (yanlış seçimin telafisi).

    Mazeret sınavı onaylanmışsa (yapıldıysa) REDDEDİLİR: o andan itibaren bağ,
    sınavın kime yapıldığının kaydıdır. Dağıtılmış oturumda yerleşim kendiliğinden
    DEĞİŞMEZ — fark oturum sayfasındaki "yeniden dağıtın" bandına döner. Kayıt
    "Mazeretli" kalır ve yeniden "mazeret sınavı bekleyen" olur.
    """
    temiz = list(dict.fromkeys(int(i) for i in record_ids))
    if not temiz:
        raise ValidationError("Mazeret sınavından çıkarılacak öğrenci seçin.")
    records = list(
        ExamAttendanceRecord.objects.filter(pk__in=temiz).select_related(
            "makeup_course", "makeup_course__session"
        )
    )
    if len(records) != len(temiz):
        raise ValidationError("Seçilen kayıtlardan biri bulunamadı; listeyi yenileyin.")
    for r in records:
        sc = live_makeup_course(r)
        if sc is None:
            raise ValidationError(f"Okul No {r.student_number}: mazeret sınavına alınmamış.")
        if sc.session.status in (ExamSessionStatus.APPROVED, ExamSessionStatus.ARCHIVED):
            raise ValidationError(
                f"Okul No {r.student_number}: '{sc.session.name}' onaylanmış; yapılmış mazeret "
                "sınavından öğrenci çıkarılamaz."
            )
    for r in records:
        r.makeup_course = None
        r.save(update_fields=["makeup_course", "updated_at"])
    return len(records)


# ---------------------------------------------------------------------------
# Mazeret takip raporu — PDF (resmî antetli) + Excel (19.09.2026, kullanıcı kararı)
# ---------------------------------------------------------------------------
def _tr_date(iso: str | None) -> str:
    return date.fromisoformat(iso).strftime("%d.%m.%Y") if iso else ""


def _report_rows(semester_id: int) -> list[dict[str, Any]]:
    """Rapor satırları: ekran satırı + gg.aa.yyyy tarihler (şablon ISO biçimlemez)."""
    return [
        {
            **row,
            "exam_date_tr": _tr_date(row["exam_date"]),
            "deadline_tr": _tr_date(row["notice_deadline"]),
            "makeup_date_tr": _tr_date(row["makeup_date"]),
            # Excel'de tek sözcük ("Ülke") yetmez; PDF'teki ifadeyle aynı olsun.
            "session_scope_label": f"{row['session_type_label']} geneli",
        }
        for row in absence_rows(semester_id)
    ]


def makeup_report_context(semester_id: int) -> dict[str, Any]:
    """PDF ve Excel'in ortak bağlamı — dört bölüm (modül başındaki kurallar)."""
    from apps.okul.models import SchoolConfig
    from shared.letterhead import letterhead_context

    donem = okul_selectors.get_school_term(semester_id)
    if donem is None:
        raise ValidationError({"semester": "Dönem bulunamadı."})
    config = SchoolConfig.load()
    rows = _report_rows(semester_id)
    return {
        **letterhead_context(
            school_name=config.school_name,
            unit="Okul Müdürlüğü",
            district=config.district,
            principal_name=config.principal_name,
        ),
        "year_label": str(donem.school_year.name),
        "term_sequence": donem.sequence,
        "term_label": str(donem),
        "generated_at": timezone.localtime().strftime("%d.%m.%Y %H:%M"),
        "summary": absence_summary(rows),
        "rows": rows,
        "g_rows": [r for r in rows if r["excuse_status"] == ExcuseStatus.UNEXCUSED],
        "waiting_rows": [r for r in rows if r["can_makeup"]],
        "external_rows": [
            r for r in rows if r["external"] and r["excuse_status"] == ExcuseStatus.EXCUSED
        ],
        "principal_name": config.principal_name,
    }


def render_makeup_report_pdf(semester_id: int) -> bytes:
    from django.template.loader import render_to_string

    from shared.pdf import html_to_pdf

    return html_to_pdf(
        render_to_string("sinav/mazeret_takip.html", makeup_report_context(semester_id))
    )


def build_makeup_workbook(semester_id: int) -> bytes:
    """Excel çalışma kopyası: dört sayfa, PDF'le aynı bölümler (e-Okul'a işleme için)."""
    import io

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    ctx = makeup_report_context(semester_id)
    wb = Workbook()
    sayfalar: list[tuple[str, str, list[tuple[str, str, int]], list[dict[str, Any]]]] = [
        (
            "Girmeyenler",
            "Sınava girmeyen öğrenciler",
            [
                ("Sınav tarihi", "exam_date_tr", 12),
                ("Sınav", "course_label", 34),
                ("Oturum", "session_name", 24),
                ("Okul No", "student_number", 9),
                ("Adı Soyadı", "full_name", 26),
                ("Sınıf/Şube", "class_label", 10),
                ("Mazeret durumu", "excuse_label", 14),
                ("Belge / açıklama", "note", 30),
                ("Bildirim son günü", "deadline_tr", 14),
                ("Mazeret sınavı", "makeup_session_name", 22),
                ("Mazeret sınavı tarihi", "makeup_date_tr", 14),
                ("Sonuç", "makeup_result_label", 10),
            ],
            ctx["rows"],
        ),
        (
            "e-Okul G",
            "e-Okul'a 'G' işlenecekler (Yönerge md. 6/1-f)",
            [
                ("Okul No", "student_number", 9),
                ("Adı Soyadı", "full_name", 26),
                ("Sınıf/Şube", "class_label", 10),
                ("Sınav", "course_label", 34),
                ("Sınav tarihi", "exam_date_tr", 12),
                ("Oturum", "session_name", 24),
            ],
            ctx["g_rows"],
        ),
        (
            "Mazeret bekleyenler",
            "Mazeret sınavı bekleyenler (Mazeretli, henüz oturuma alınmadı)",
            [
                ("Okul No", "student_number", 9),
                ("Adı Soyadı", "full_name", 26),
                ("Sınıf/Şube", "class_label", 10),
                ("Sınav", "course_label", 34),
                ("Sınav tarihi", "exam_date_tr", 12),
            ],
            ctx["waiting_rows"],
        ),
        (
            "İl-İlçe bildirimi",
            "İl/ilçe MEM'e bildirilecekler — ülke/il/ilçe geneli sınavlar (Yönerge md. 5/1-z)",
            [
                ("Okul No", "student_number", 9),
                ("Adı Soyadı", "full_name", 26),
                ("Sınıf/Şube", "class_label", 10),
                ("Sınav", "course_label", 34),
                ("Sınav türü", "session_scope_label", 14),
                ("Sınav tarihi", "exam_date_tr", 12),
            ],
            ctx["external_rows"],
        ),
    ]
    for sira, (ad, baslik, sutunlar, satirlar) in enumerate(sayfalar):
        ws = wb.active if sira == 0 else wb.create_sheet()
        assert ws is not None
        ws.title = ad
        ws.cell(row=1, column=1, value=f"{ctx['school_name']} — {baslik}").font = Font(
            bold=True, size=12
        )
        ws.cell(row=2, column=1, value=f"{ctx['term_label']} · {ctx['generated_at']}")
        for col, (etiket, _alan, genislik) in enumerate(sutunlar, start=1):
            hucre = ws.cell(row=4, column=col, value=etiket)
            hucre.font = Font(bold=True)
            hucre.fill = PatternFill("solid", fgColor="D9D9D9")
            hucre.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(col)].width = genislik
        for satir_no, satir in enumerate(satirlar, start=5):
            for col, (_etiket, alan, _g) in enumerate(sutunlar, start=1):
                ws.cell(row=satir_no, column=col, value=satir.get(alan) or "")
        ws.freeze_panes = "A5"
    tampon = io.BytesIO()
    wb.save(tampon)
    return tampon.getvalue()


#: Takip raporu biçimleri — `?format=` DRF'nin içerik müzakeresine ayrılmıştır, uç `kind` alır.
REPORT_KINDS = ("pdf", "xlsx")


def render_makeup_report(semester_id: int, kind: str) -> ReportFile:
    """İndirme dosyası — ad ASCII ve dönem kimlikli (oturum evrakıyla aynı desen)."""
    from apps.sinav.services import _PDF_MIME, _XLSX_MIME, ReportFile

    if kind not in REPORT_KINDS:
        raise ValidationError({"kind": "Rapor biçimi pdf ya da xlsx olmalı."})
    if kind == "xlsx":
        return ReportFile(
            filename=f"mazeret_takip_donem_{semester_id}.xlsx",
            content_type=_XLSX_MIME,
            content=build_makeup_workbook(semester_id),
        )
    return ReportFile(
        filename=f"mazeret_takip_donem_{semester_id}.pdf",
        content_type=_PDF_MIME,
        content=render_makeup_report_pdf(semester_id),
    )
